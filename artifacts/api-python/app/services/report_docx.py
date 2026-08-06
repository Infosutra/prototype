"""DOCX export for DQA Daily & Final reports."""

from __future__ import annotations

from io import BytesIO
from typing import Any, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor


TEAL = RGBColor(0x13, 0x4E, 0x4A)
MUTED = RGBColor(0x57, 0x53, 0x4E)
CAPTION = RGBColor(0x78, 0x71, 0x6C)


def _set_run_font(run, *, size_pt: float = 10, bold: bool = False, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def _add_paragraph(
    doc: Document,
    text: str,
    *,
    size_pt: float = 10,
    bold: bool = False,
    color: RGBColor | None = None,
    space_after: float = 6,
    space_before: float = 0,
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    run = p.add_run(text or "")
    _set_run_font(run, size_pt=size_pt, bold=bold, color=color)


def _add_heading(doc: Document, text: str, *, level: int = 1) -> None:
    if level <= 1:
        _add_paragraph(doc, text, size_pt=16, bold=True, space_after=8, space_before=4)
    elif level == 2:
        _add_paragraph(doc, text, size_pt=12, bold=True, color=TEAL, space_after=6, space_before=14)
    else:
        _add_paragraph(doc, text, size_pt=11, bold=True, color=TEAL, space_after=4, space_before=10)


def _add_paras(doc: Document, text: str) -> None:
    for part in (text or "").split("\n\n"):
        part = part.strip()
        if part:
            _add_paragraph(doc, part, size_pt=10, space_after=6)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
) -> None:
    if not headers:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, size_pt=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade_cell(cell, "134E4A")
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = "" if value is None else str(value)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, size_pt=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _add_chart(doc: Document, png: bytes | None, caption: str, *, width_in: float = 6.2) -> None:
    if not png:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(BytesIO(png), width=Inches(width_in))
    _add_paragraph(doc, caption, size_pt=8, color=CAPTION, space_after=10)


def _save(doc: Document) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_daily_docx(stats: dict[str, Any]) -> bytes:
    from app.services.dqa_daily_report import _daily_chart_pngs

    charts = _daily_chart_pngs(stats)
    tot = stats["totals"]
    day = stats.get("dayNumber")
    day_label = f"Day {day}" if day is not None else "Day n/a"

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    _add_paragraph(
        doc,
        stats.get("organizationName") or "Infosutra",
        size_pt=9,
        color=MUTED,
        space_after=2,
    )
    _add_heading(doc, f"DQA Daily Report — {stats['studyName']}", level=1)
    _add_paragraph(
        doc,
        (
            f"Report: Daily DQA — {day_label} · "
            f"{stats.get('reportDateDisplay') or stats['reportDate']} "
            f"({stats.get('timezone') or 'Asia/Kolkata'}) · "
            f"KoBo pull {stats.get('lastKoboPullDisplay') or 'never'} · "
            f"Generated {stats.get('generatedAtDisplay') or 'n/a'} · "
            f"New {tot['newToday']} · Cumulative {tot['cumulative']}"
        ),
        size_pt=8,
        color=MUTED,
        space_after=10,
    )

    _add_paragraph(doc, "Today's headline", size_pt=10, bold=True, space_after=4)
    _add_paras(doc, stats.get("aiHeadline") or "")

    _add_heading(doc, "1.1 Today's intake & flags — by tool", level=2)
    tool_rows = [
        [
            t["toolCode"],
            t["newToday"],
            t["cumulative"],
            t["redToday"],
            t["amberToday"],
            f"{t['flaggedPctToday']}%",
        ]
        for t in stats.get("tools") or []
    ]
    _add_table(
        doc,
        ["Tool", "New today", "Cumulative", "RED", "AMBER", "% flagged"],
        tool_rows,
    )
    _add_paragraph(
        doc,
        "Note: today's flag rate is inflated by same-day records not yet back-checked; "
        "the cumulative trend (Section 1.4) is the truer picture.",
        size_pt=9,
        color=MUTED,
    )
    _add_chart(doc, charts.get("flagsToday"), "Figure — Today's RED / AMBER by tool.")
    _add_chart(doc, charts.get("topRules"), "Figure — Top failing rules today.", width_in=5.6)

    _add_heading(doc, "1.2 RED items to correct or back-check (priority)", level=2)
    red_rows = [
        [
            r["toolCode"],
            r["ruleId"],
            (r["title"] or "")[:60],
            r["count"],
            f"{r.get('exampleEnumerator') or '—'} · UDISE {r.get('exampleUdise') or '—'}",
        ]
        for r in (stats.get("redGrouped") or [])[:20]
    ]
    _add_table(doc, ["Tool", "Rule", "Check", "Records", "Example"], red_rows)

    _add_heading(doc, "1.3 Enumerators to back-check tomorrow", level=2)
    enum_rows = [
        [
            e["enumerator"],
            ", ".join(e.get("tools") or []) or "—",
            e["submissionsToday"],
            f"{e['flagRate']}%",
            e.get("medianTimeLabel") or "—",
            e.get("whyFlagged") or "—",
        ]
        for e in (stats.get("enumerators") or [])[:12]
    ]
    _add_table(
        doc,
        ["Enumerator", "Tool(s)", "Records", "Flag %", "Median", "Why flagged"],
        enum_rows,
    )

    _add_heading(doc, "1.4 Coverage & trend", level=2)
    _add_paras(doc, stats.get("aiCoverageNote") or "")
    _add_chart(doc, charts.get("coverage"), "Figure — Cumulative submissions vs study targets.")
    _add_chart(doc, charts.get("trend"), "Figure — Cumulative flag rate by study day.")

    _add_heading(doc, "Sign-off checklist (read-only)", level=2)
    for item in stats.get("signOffChecklist") or []:
        _add_paragraph(
            doc,
            f"[{(item.get('severity') or '').upper()}] {item.get('label') or ''} — {item.get('detail') or ''}",
            size_pt=9,
            space_after=3,
        )
    _add_paragraph(
        doc,
        "Infosutra · Kobo is source of truth (read-only) · Checklist is informational.",
        size_pt=8,
        color=MUTED,
        space_before=10,
    )
    return _save(doc)


def render_final_docx(stats: dict[str, Any]) -> bytes:
    from app.services.dqa_final_report import _final_chart_pngs, _section_prose

    charts = _final_chart_pngs(stats)
    prose = stats.get("sectionProse") or _section_prose(stats)
    tot = stats["totals"]
    tools = stats.get("tools") or []

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    _add_paragraph(
        doc,
        stats.get("organizationName") or "Infosutra",
        size_pt=9,
        color=MUTED,
        space_after=2,
    )
    _add_heading(doc, stats.get("title") or "Final DQA", level=1)
    _add_paragraph(
        doc,
        (
            f"Final DQA — end of round · "
            f"{stats.get('reportDateDisplay') or stats['reportDate']} · "
            f"KoBo pull {stats.get('lastKoboPullDisplay') or 'n/a'} · "
            f"Generated {stats.get('generatedAtDisplay') or 'n/a'} · "
            f"Total submissions {tot['cumulative']}"
        ),
        size_pt=8,
        color=MUTED,
        space_after=10,
    )

    _add_heading(doc, "2.1 Executive summary", level=2)
    _add_paras(doc, stats.get("aiHeadline") or "")

    _add_heading(doc, "2.2 Coverage — by tool", level=2)
    _add_paras(doc, prose.get("coverage") or "")
    _add_table(
        doc,
        ["Tool", "Target", "Actual", "% of plan"],
        [
            [
                t["toolCode"],
                t["target"] if t.get("target") is not None else "—",
                t["cumulative"],
                f"{t['coveragePct']}%" if t.get("coveragePct") is not None else "—",
            ]
            for t in tools
        ],
    )
    _add_chart(doc, charts.get("coverage"), "Figure — Actual vs plan by tool.")
    _add_paras(doc, prose.get("trend") or "")
    _add_chart(doc, charts.get("trend"), "Figure — Cumulative flag rate across the study window.")

    _add_heading(doc, "2.3 Data-quality flag summary — by tool", level=2)
    _add_paras(doc, prose.get("flagSummary") or "")
    flagged_total = sum(int(t.get("flaggedSubmissions") or 0) for t in tools)
    flag_rows = [
        [
            t["toolCode"],
            t["cumulative"],
            t["redCumulative"],
            t["amberCumulative"],
            f"{t['flaggedPct']}%",
        ]
        for t in tools
    ]
    flag_rows.append(
        [
            "Overall",
            tot["cumulative"],
            tot["redOpen"],
            tot["amberOpen"],
            f"{round(100.0 * flagged_total / tot['cumulative'], 1) if tot.get('cumulative') else 0}%",
        ]
    )
    _add_table(doc, ["Tool", "Submissions", "RED", "AMBER", "% flagged"], flag_rows)
    _add_chart(
        doc,
        charts.get("flagSummary"),
        "Figure 1 — DQA flag summary by tool. Submissions, RED/AMBER counts and % flagged.",
    )
    _add_chart(doc, charts.get("topRules"), "Figure — Top failing rules (cumulative).", width_in=5.6)

    _add_heading(doc, "2.4 Findings by tool", level=2)
    for block in stats.get("findingsByTool") or []:
        _add_heading(doc, f"{block.get('toolCode')} — {block.get('projectName')}", level=3)
        _add_paragraph(doc, block.get("summary") or "", size_pt=10, space_after=4)
        _add_table(
            doc,
            ["Rule", "Check", "Records", "Sev"],
            [
                [
                    r.get("ruleId"),
                    (r.get("title") or "")[:70],
                    r.get("count"),
                    (r.get("severity") or "").upper(),
                ]
                for r in block.get("rules") or []
            ],
        )

    _add_heading(doc, "2.5 Enumerator performance", level=2)
    _add_paragraph(
        doc,
        "Enumerators ranked by RED volume and flag rate. Median time marked “(below)” when under the group median.",
        size_pt=9,
        color=MUTED,
    )
    _add_table(
        doc,
        ["Enumerator", "Tool(s)", "Records", "Flag %", "Median", "Action"],
        [
            [
                e["enumerator"],
                ", ".join(e.get("tools") or []) or "—",
                e["submissions"],
                f"{e['flagRate']}%",
                e.get("medianTimeLabel") or "—",
                e.get("action") or "—",
            ]
            for e in (stats.get("enumeratorsAll") or [])[:12]
        ],
    )

    _add_heading(doc, "2.6 Triangulation across tools", level=2)
    _add_paras(doc, prose.get("tr1") or "")
    _add_chart(
        doc,
        charts.get("tr1"),
        "Figure 2 — TR-1 teacher practice, claimed vs observed. Say–do gap per practice.",
    )
    _add_paras(doc, prose.get("triangulationFurther") or "")
    _add_chart(doc, charts.get("tr3"), "Figure — TR-3 governance concordance.")

    _add_heading(doc, "2.7 What remains before sign-off", level=2)
    _add_table(
        doc,
        ["Action", "Tool", "Records", "Owner", "Sev"],
        [
            [
                (c.get("label") or "")[:60],
                c.get("toolCode") or "—",
                c.get("records") if c.get("records") is not None else "—",
                c.get("owner") or c.get("ownerHint") or "—",
                (c.get("severity") or "").upper(),
            ]
            for c in stats.get("signOffChecklist") or []
        ],
    )
    _add_paras(doc, prose.get("signOffClose") or "")
    _add_paragraph(
        doc,
        "Infosutra Final DQA · Kobo is source of truth (read-only) · Checklist is informational.",
        size_pt=8,
        color=MUTED,
        space_before=10,
    )
    return _save(doc)


def render_report_docx(report_type: str | None, stats: dict[str, Any]) -> bytes:
    if report_type == "final_dqa":
        return render_final_docx(stats)
    return render_daily_docx(stats)
