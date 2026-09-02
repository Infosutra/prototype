"""DOCX rendering for diarized audio transcripts."""

from __future__ import annotations

from docx import Document
from docx.shared import Pt, RGBColor

from app.rendering.docx_helpers import _save, _set_run_font, _shade_cell
from app.services.transcript_export import TranscriptExportContext, format_duration_minutes, format_timestamp

NAVY = RGBColor(0x1A, 0x36, 0x5D)
LABEL_FILL = "E8F0FE"


def _add_title(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    _set_run_font(run, size_pt=18, bold=True, color=NAVY)


def _add_subtitle(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    _set_run_font(run, size_pt=11, bold=True, color=NAVY)


def _add_metadata_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for row_index, (label, value) in enumerate(rows):
        label_cell = table.rows[row_index].cells[0]
        value_cell = table.rows[row_index].cells[1]
        label_cell.text = label
        value_cell.text = value
        for paragraph in label_cell.paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, size_pt=9, bold=True, color=NAVY)
        _shade_cell(label_cell, LABEL_FILL)
        for paragraph in value_cell.paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, size_pt=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def _add_section_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    _set_run_font(run, size_pt=12, bold=True, color=NAVY)


def _metadata_rows(context: TranscriptExportContext) -> list[tuple[str, str]]:
    provider = (context.transcription_provider or "—").upper()
    if context.transcription_model:
        provider = f"{provider} ({context.transcription_model})"

    cost = "—"
    if context.transcription_cost_amount is not None and context.transcription_cost_amount > 0:
        cost = (
            f"{context.transcription_cost_currency or 'INR'} "
            f"{context.transcription_cost_amount:.2f}"
        )

    rows = [
        ("Source File", context.original_filename),
        ("Total Utterances", f"{context.total_utterances:,}"),
        ("Speakers", context.speaker_stats),
        ("Study", context.study_name),
        ("Transcription Provider", provider),
        ("Language", context.transcription_language or "—"),
        ("Words", f"{context.word_count:,}"),
        ("Transcription Cost", cost),
        ("Transcribed", context.transcribed_at or "—"),
        ("Exported", context.exported_at),
    ]
    if context.description.strip():
        rows.insert(3, ("Description", context.description.strip()))
    return rows


def render_transcript_docx(context: TranscriptExportContext) -> bytes:
    doc = Document()
    title = f"Transcript — {context.recording_name}"
    subtitle = (
        f"{context.study_name} | {context.transcribed_at or '—'} | "
        f"{format_duration_minutes(context.duration_seconds)}"
    )

    _add_title(doc, title)
    _add_subtitle(doc, subtitle)
    _add_metadata_table(doc, _metadata_rows(context))
    _add_section_heading(doc, "Diarized transcript")

    for segment in context.segments:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        time_range = (
            f"[{format_timestamp(segment['start_seconds'])} – "
            f"{format_timestamp(segment['end_seconds'])}]"
        )
        prefix = f"{time_range} {segment['speaker']}: "
        prefix_run = paragraph.add_run(prefix)
        _set_run_font(prefix_run, size_pt=9, bold=True, color=NAVY)
        text_run = paragraph.add_run(segment["text"])
        _set_run_font(text_run, size_pt=9)

    return _save(doc)
