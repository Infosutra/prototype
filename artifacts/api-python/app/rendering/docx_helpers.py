"""Shared DOCX helpers for DQA reports."""

from __future__ import annotations

from io import BytesIO
from typing import Any, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor

TEAL = RGBColor(0x13, 0x4E, 0x4A)
MUTED = RGBColor(0x57, 0x53, 0x4E)
CAPTION = RGBColor(0x78, 0x71, 0x6C)

def _set_run_font(run, *, size_pt: float = 10, bold: bool = False, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def _add_paragraph(
    doc: Document,
    text: str,
    *,
    size_pt: float = 10,
    bold: bool = False,
    color: RGBColor | None = None,
    space_after: float = 6,
    space_before: float = 0,
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    run = p.add_run(text or "")
    _set_run_font(run, size_pt=size_pt, bold=bold, color=color)


def _add_heading(doc: Document, text: str, *, level: int = 1) -> None:
    if level <= 1:
        _add_paragraph(doc, text, size_pt=16, bold=True, space_after=8, space_before=4)
    elif level == 2:
        _add_paragraph(doc, text, size_pt=12, bold=True, color=TEAL, space_after=6, space_before=14)
    else:
        _add_paragraph(doc, text, size_pt=11, bold=True, color=TEAL, space_after=4, space_before=10)


def _add_paras(doc: Document, text: str) -> None:
    for part in (text or "").split("\n\n"):
        part = part.strip()
        if part:
            _add_paragraph(doc, part, size_pt=10, space_after=6)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
) -> None:
    if not headers:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, size_pt=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade_cell(cell, "134E4A")
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = "" if value is None else str(value)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, size_pt=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _add_chart(doc: Document, png: bytes | None, caption: str, *, width_in: float = 6.2) -> None:
    if not png:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(BytesIO(png), width=Inches(width_in))
    _add_paragraph(doc, caption, size_pt=8, color=CAPTION, space_after=10)


def _save(doc: Document) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
