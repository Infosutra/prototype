"""DOCX export for DQA Daily reports."""

from __future__ import annotations

from typing import Any

from docx import Document
from docx.shared import Cm

from app.rendering.daily import daily_chart_pngs
from app.rendering.docx_helpers import (
    MUTED,
    TEAL,
    _add_chart,
    _add_heading,
    _add_paragraph,
    _add_paras,
    _add_table,
    _save,
)

def render_daily_docx(stats: dict[str, Any]) -> bytes:
    charts = daily_chart_pngs(stats)
    tot = stats["totals"]
    day = stats.get("dayNumber")
    day_label = f"Day {day}" if day is not None else "Day n/a"

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    _add_paragraph(
        doc,
        stats.get("organizationName") or "Infosutra",
        size_pt=9,
        color=MUTED,
        space_after=2,
    )
    _add_heading(doc, f"DQA Daily Report — {stats['studyName']}", level=1)
    _add_paragraph(
        doc,
        (
            f"Report: Daily DQA — {day_label} · "
            f"{stats.get('reportDateDisplay') or stats['reportDate']} "
            f"({stats.get('timezone') or 'Asia/Kolkata'}) · "
            f"KoBo pull {stats.get('lastKoboPullDisplay') or 'never'} · "
            f"Generated {stats.get('generatedAtDisplay') or 'n/a'} · "
            f"New {tot['newToday']} · Cumulative {tot['cumulative']}"
        ),
        size_pt=8,
        color=MUTED,
        space_after=10,
    )

    _add_paragraph(doc, "Today's headline", size_pt=10, bold=True, space_after=4)
    _add_paras(doc, stats.get("aiHeadline") or "")

    _add_heading(doc, "1.1 Today's intake & flags — by tool", level=2)
    tool_rows = [
        [
            t["toolCode"],
            t["newToday"],
            t["cumulative"],
            t["redToday"],
            t["amberToday"],
            f"{t['flaggedPctToday']}%",
        ]
        for t in stats.get("tools") or []
    ]
    _add_table(
        doc,
        ["Tool", "New today", "Cumulative", "RED", "AMBER", "% flagged"],
        tool_rows,
    )
    _add_paragraph(
        doc,
        "Note: today's flag rate is inflated by same-day records not yet back-checked; "
        "the cumulative trend (Section 1.4) is the truer picture.",
        size_pt=9,
        color=MUTED,
    )
    _add_chart(doc, charts.get("flagsToday"), "Figure — Today's RED / AMBER by tool.")
    _add_chart(doc, charts.get("topRules"), "Figure — Top failing rules today.", width_in=5.6)

    _add_heading(doc, "1.2 RED items to correct or back-check (priority)", level=2)
    red_rows = [
        [
            r["toolCode"],
            r["ruleId"],
            (r["title"] or "")[:60],
            r["count"],
            f"{r.get('exampleEnumerator') or '—'} · UDISE {r.get('exampleUdise') or '—'}",
        ]
        for r in (stats.get("redGrouped") or [])[:20]
    ]
    _add_table(doc, ["Tool", "Rule", "Check", "Records", "Example"], red_rows)

    _add_heading(doc, "1.3 Enumerators to back-check tomorrow", level=2)
    enum_rows = [
        [
            e["enumerator"],
            ", ".join(e.get("tools") or []) or "—",
            e["submissionsToday"],
            f"{e['flagRate']}%",
            e.get("medianTimeLabel") or "—",
            e.get("whyFlagged") or "—",
        ]
        for e in (stats.get("enumerators") or [])[:12]
    ]
    _add_table(
        doc,
        ["Enumerator", "Tool(s)", "Records", "Flag %", "Median", "Why flagged"],
        enum_rows,
    )

    _add_heading(doc, "1.4 Coverage & trend", level=2)
    _add_paras(doc, stats.get("aiCoverageNote") or "")
    _add_chart(doc, charts.get("coverage"), "Figure — Cumulative submissions vs study targets.")
    _add_chart(doc, charts.get("trend"), "Figure — Cumulative flag rate by study day.")

    _add_heading(doc, "Sign-off checklist (read-only)", level=2)
    for item in stats.get("signOffChecklist") or []:
        _add_paragraph(
            doc,
            f"[{(item.get('severity') or '').upper()}] {item.get('label') or ''} — {item.get('detail') or ''}",
            size_pt=9,
            space_after=3,
        )
    _add_paragraph(
        doc,
        "Infosutra · Kobo is source of truth (read-only) · Checklist is informational.",
        size_pt=8,
        color=MUTED,
        space_before=10,
    )
    return _save(doc)
