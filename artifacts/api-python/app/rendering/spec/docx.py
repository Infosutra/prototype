"""DOCX rendering of a Report Specification."""

from __future__ import annotations

from typing import Any

from app.domain.report_spec.spec import (
    BarChartComponent,
    KpiGroupComponent,
    LineChartComponent,
    MetricComponent,
    NarrativeComponent,
    PieChartComponent,
    ProgressComponent,
    RankingComponent,
    StackedBarChartComponent,
    TableComponent,
    TextComponent,
)
from app.rendering.docx_helpers import (
    CAPTION,
    MUTED,
    _add_chart,
    _add_heading,
    _add_paragraph,
    _add_paras,
    _add_table,
    _save,
)
from app.rendering.spec import charts as spec_charts
from app.rendering.spec.base import RenderPayload, format_value, highlight_matches

_CALLOUT_LABEL = {"insight": "Insight", "warning": "Warning", "action_plan": "Action plan"}


def _unavailable(doc, reason: str) -> None:
    _add_paragraph(doc, reason, size_pt=9, color=CAPTION, space_after=8)


def _component(doc, component: Any, payload: RenderPayload) -> None:
    if isinstance(component, NarrativeComponent):
        text = (payload.narratives.get(component.id) or "").strip()
        if not text:
            return
        _add_paragraph(
            doc,
            component.title or _CALLOUT_LABEL[component.type],
            size_pt=10,
            bold=True,
            space_after=4,
        )
        _add_paras(doc, text)
        return

    if component.title:
        _add_heading(doc, component.title, level=3)

    if isinstance(component, TextComponent):
        if component.style == "caption":
            _add_paragraph(doc, component.body, size_pt=8, color=CAPTION, space_after=8)
        else:
            _add_paras(doc, component.body)
        return

    if isinstance(component, MetricComponent):
        record, error = payload.record_for(component)
        if error:
            _unavailable(doc, error)
            return
        value = format_value(record.get(component.field), component.format, component.unit)
        _add_paragraph(doc, f"{component.label}: {value}", size_pt=10, space_after=4)
        if component.caption:
            _add_paragraph(doc, component.caption, size_pt=8, color=CAPTION, space_after=8)
        return

    if isinstance(component, KpiGroupComponent):
        record, error = payload.record_for(component)
        if error:
            _unavailable(doc, error)
            return
        headers = [item.label for item in component.items]
        row = [
            format_value(record.get(item.field), item.format, item.unit)
            for item in component.items
        ]
        _add_table(doc, headers, [row])
        return

    if isinstance(component, TableComponent):
        rows, error = payload.rows_for(component)
        if error:
            _unavailable(doc, error)
            return
        if not rows:
            _add_paragraph(doc, component.empty_text, size_pt=9, color=MUTED, space_after=8)
            return
        headers = [col.label for col in component.columns]
        body = [
            [format_value(row.get(col.field), col.format) for col in component.columns]
            for row in rows
        ]
        _add_table(doc, headers, body)
        return

    if isinstance(component, RankingComponent):
        rows, error = payload.rows_for(component)
        if error:
            _unavailable(doc, error)
            return
        ordered = sorted(
            rows,
            key=lambda row: float(row.get(component.value_field) or 0),
            reverse=component.order == "desc",
        )[: component.limit]
        for idx, row in enumerate(ordered, start=1):
            suffix = ""
            if highlight_matches(row, component.highlight) and component.highlight:
                suffix = f" — {component.highlight.label}"
            value = format_value(row.get(component.value_field), component.value_format)
            _add_paragraph(
                doc,
                f"{idx}. {row.get(component.label_field) or '—'} — {value}{suffix}",
                size_pt=10,
                space_after=2,
            )
        return

    rows, error = payload.rows_for(component)
    if error:
        _unavailable(doc, error)
        return

    png: bytes | None = None
    if isinstance(component, StackedBarChartComponent):
        png = spec_charts.bar_chart(
            rows,
            x=component.x,
            series=component.series,
            title=component.title,
            y_label=component.y_label,
            stacked=True,
        )
    elif isinstance(component, BarChartComponent):
        png = spec_charts.bar_chart(
            rows,
            x=component.x,
            series=component.series,
            title=component.title,
            y_label=component.y_label,
            orientation=component.orientation,
        )
    elif isinstance(component, LineChartComponent):
        png = spec_charts.line_chart(
            rows,
            x=component.x,
            series=component.series,
            title=component.title,
            y_label=component.y_label,
            area=component.area,
        )
    elif isinstance(component, PieChartComponent):
        png = spec_charts.pie_chart(
            rows[: component.limit],
            label_field=component.label_field,
            value_field=component.value_field,
            title=component.title,
            donut=component.donut,
        )
    elif isinstance(component, ProgressComponent):
        png = spec_charts.progress_chart(
            rows[: component.limit],
            label_field=component.label_field,
            value_field=component.value_field,
            target_field=component.target_field,
            title=component.title,
        )

    if png:
        _add_chart(doc, png, getattr(component, "caption", None) or "")


def render_spec_docx(payload: RenderPayload) -> bytes:
    from docx import Document

    spec = payload.spec
    doc = Document()
    _add_paragraph(
        doc,
        payload.meta.get("organizationName") or "Infosutra",
        size_pt=9,
        color=MUTED,
        space_after=2,
    )
    _add_heading(doc, spec.title, level=1)
    if spec.subtitle:
        _add_paragraph(doc, spec.subtitle, size_pt=10, color=MUTED, space_after=6)
    for label, value in payload.meta.get("rows", []):
        if value:
            _add_paragraph(doc, f"{label}: {value}", size_pt=9, color=MUTED, space_after=2)

    for section in spec.sections:
        _add_heading(doc, section.title, level=2)
        if section.description:
            _add_paras(doc, section.description)
        for component in section.components:
            _component(doc, component, payload)

    footer = payload.meta.get("footer")
    if footer:
        _add_paragraph(doc, footer, size_pt=8, color=CAPTION, space_before=12)
    return _save(doc)
